"""원본 Word 템플릿과 결과 문서의 형식 속성을 비교해 레이아웃 붕괴를 검출한다.

1) 형식이 있는 PAD 스타일 템플릿을 만든다 (맑은 고딕 10pt, numbering, 머리글/바닥글, 들여쓰기)
2) 'LLM이 텍스트만 보고 다시 만든 문서'를 재현한다 (형식 정보 소실)
3) 두 문서의 형식 속성을 비교한다

pip install python-docx
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).parent


def set_font(run, name: str, size: float) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def make_template(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    set_font(sec.header.paragraphs[0].add_run("MID · Part Approval Document"), "맑은 고딕", 9)
    set_font(sec.footer.paragraphs[0].add_run("Confidential"), "맑은 고딕", 8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(title.add_run("1. 부품 정보"), "맑은 고딕", 12)

    for text in ["부품 번호: {{ part_number }}", "제조사: {{ manufacturer }}"]:
        p = doc.add_paragraph(style="List Number")
        set_font(p.add_run(text), "맑은 고딕", 10)

    body = doc.add_paragraph()
    body.paragraph_format.left_indent = Pt(20)
    set_font(body.add_run("시험 요약: {{ test_summary }}"), "맑은 고딕", 10)
    doc.save(path)


def make_llm_regenerated(template: Path, path: Path) -> None:
    """LLM이 템플릿의 '텍스트'만 보고 내용을 바꿔 다시 만든 문서 (실제로 겪은 증상 재현)."""
    src = Document(template)
    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].add_run("MID · Part Approval Document")       # 머리글 스타일 소실
    sec.footer.paragraphs[0].add_run("Confidential")
    sec.footer.add_paragraph("Confidential")                               # 바닥글 중복
    values = {"{{ part_number }}": "PN-2026-0412", "{{ manufacturer }}": "ACME Space",
              "{{ test_summary }}": "열진공 시험 통과"}
    for i, para in enumerate(src.paragraphs):
        text = para.text
        for k, v in values.items():
            text = text.replace(k, v)
        if para.style.name == "List Number":
            text = f"{i}. {text}"                                          # numbering → 일반 텍스트
        p = doc.add_paragraph()                                            # 들여쓰기 소실
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER                        # 불필요한 중앙 정렬
        set_font(p.add_run(text), "Noto Sans CJK KR", 10.5)                # 폰트·크기 변경
    doc.save(path)


def snapshot(path: Path) -> dict:
    doc = Document(path)
    sec = doc.sections[0]
    paras = []
    for p in doc.paragraphs:
        run = p.runs[0] if p.runs else None
        paras.append({
            "text": p.text[:18],
            "style": p.style.name,
            "align": str(p.alignment).split(".")[-1].split(" ")[0] if p.alignment is not None else "INHERIT",
            "indent": p.paragraph_format.left_indent.pt if p.paragraph_format.left_indent else 0,
            "font": run.font.name if run else None,
            "size": run.font.size.pt if run and run.font.size else None,
        })
    return {
        "header": [p.text for p in sec.header.paragraphs if p.text],
        "footer": [p.text for p in sec.footer.paragraphs if p.text],
        "paragraphs": paras,
    }


def diff(a: dict, b: dict) -> list[str]:
    issues = []
    if a["header"] != b["header"]:
        issues.append(f"머리글 변경: {a['header']} → {b['header']}")
    if len(b["footer"]) > len(a["footer"]):
        issues.append(f"바닥글 중복: {len(a['footer'])}개 → {len(b['footer'])}개")
    for i, (pa, pb) in enumerate(zip(a["paragraphs"], b["paragraphs"])):
        for key, label in [("style", "스타일(numbering)"), ("align", "정렬"), ("indent", "들여쓰기"),
                           ("font", "폰트"), ("size", "크기")]:
            if pa[key] != pb[key]:
                issues.append(f"문단{i} '{pa['text']}' {label}: {pa[key]} → {pb[key]}")
    return issues


def main() -> None:
    template, regenerated = OUT / "template_pad.docx", OUT / "llm_regenerated.docx"
    make_template(template)
    make_llm_regenerated(template, regenerated)

    issues = diff(snapshot(template), snapshot(regenerated))
    print(f"레이아웃 붕괴 {len(issues)}건 검출\n")
    for issue in issues:
        print(" ❌", issue)


if __name__ == "__main__":
    main()

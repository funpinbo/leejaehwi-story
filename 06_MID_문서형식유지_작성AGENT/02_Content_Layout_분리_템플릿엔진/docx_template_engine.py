"""Content / Layout 분리 템플릿 엔진.

- LLM은 구조화된 JSON(내용)만 만든다 → Pydantic 스키마로 검증
- 템플릿 엔진은 Word 원본의 XML·스타일을 건드리지 않고 {{ placeholder }} 텍스트만 교체한다
  · 본문, 표, 머리글, 바닥글 모두 처리
  · Word가 하나의 placeholder를 여러 run으로 쪼개 저장한 경우도 처리
    (예: "{{ part_" + "number }}" — 맞춤법 검사·편집 이력 때문에 실제로 자주 발생)
  · 템플릿에 있는데 데이터에 없는 키는 조용히 넘기지 않고 에러

pip install python-docx pydantic
"""

import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import BaseModel, Field

OUT = Path(__file__).parent
PLACEHOLDER = re.compile(r"\{\{\s*(\w+)\s*\}\}")


# ---------- Content: LLM이 만드는 JSON의 스키마 ----------
class PadContent(BaseModel):
    part_number: str = Field(pattern=r"^PN-\d{4}-\d{4}$")
    manufacturer: str
    spec_ref: str
    test_summary: str = Field(max_length=200)


# ---------- Layout: 템플릿 엔진 ----------
def iter_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def render_paragraph(paragraph, values: dict[str, str]) -> set[str]:
    runs = paragraph.runs
    texts = [r.text for r in runs]
    full = "".join(texts)
    used = set()
    if "{{" not in full:
        return used

    # 문단 전체 문자열의 위치 → (run 번호, run 내부 위치)
    owner = [(i, j) for i, t in enumerate(texts) for j in range(len(t))]
    for m in reversed(list(PLACEHOLDER.finditer(full))):  # 뒤에서부터 바꿔야 앞의 위치가 유지됨
        key = m.group(1)
        if key not in values:
            raise KeyError(f"템플릿 placeholder '{key}'에 해당하는 데이터가 없습니다")
        used.add(key)
        (rs, os_), (re_, oe) = owner[m.start()], owner[m.end() - 1]
        # placeholder가 시작한 run의 서식으로 값을 넣고, 뒤쪽 run에 걸친 조각은 지운다
        texts[rs] = texts[rs][:os_] + values[key] + (texts[rs][oe + 1 :] if rs == re_ else "")
        for k in range(rs + 1, re_):
            texts[k] = ""
        if re_ != rs:
            texts[re_] = texts[re_][oe + 1 :]
    for run, text in zip(runs, texts):
        if run.text != text:
            run.text = text  # run의 rPr(폰트·크기·굵기)은 그대로
    return used


def render(template: Path, content: BaseModel, out: Path) -> None:
    doc = Document(template)
    values = content.model_dump()
    used = set()
    for p in iter_paragraphs(doc):
        used |= render_paragraph(p, values)
    unused = set(values) - used
    if unused:
        print(f"   (참고) 템플릿에서 쓰이지 않은 데이터: {sorted(unused)}")
    doc.save(out)


# ---------- 데모 ----------
def font(run, name="맑은 고딕", size=10.0, bold=False):
    run.font.name, run.font.size, run.font.bold = name, Pt(size), bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    return run


def make_template(path: Path) -> None:
    doc = Document()
    font(doc.sections[0].header.paragraphs[0].add_run("PAD · {{ part_number }}"), size=9)
    font(doc.add_paragraph().add_run("1. 부품 정보"), size=12, bold=True)

    p = doc.add_paragraph(style="List Number")
    font(p.add_run("부품 번호: "), bold=True)
    font(p.add_run("{{ part_"))      # Word가 placeholder를 두 run으로 쪼갠 상황
    font(p.add_run("number }}"))

    p = doc.add_paragraph(style="List Number")
    font(p.add_run("제조사: "), bold=True)
    font(p.add_run("{{ manufacturer }}"))

    table = doc.add_table(rows=2, cols=2)
    for (r, c), text in {(0, 0): "적용 규격", (0, 1): "{{ spec_ref }}",
                         (1, 0): "시험 요약", (1, 1): "{{ test_summary }}"}.items():
        font(table.cell(r, c).paragraphs[0].add_run(text), bold=(c == 0))
    doc.save(path)


def fake_llm() -> str:
    """Report Agent가 만든 내용 JSON (실제로는 LLM 출력)."""
    return json.dumps({
        "part_number": "PN-2026-0412",
        "manufacturer": "ACME Space Components",
        "spec_ref": "ESCC 9000 (예시)",
        "test_summary": "열진공 시험 및 진동 시험 통과, 비정상 항목 없음",
    }, ensure_ascii=False)


def formats(path: Path) -> list[tuple]:
    doc = Document(path)
    return [(p.style.name, r.font.name, r.font.size and r.font.size.pt, r.font.bold)
            for p in iter_paragraphs(doc) for r in p.runs if r.text]


def main() -> None:
    template, out = OUT / "template_pad.docx", OUT / "rendered_pad.docx"
    make_template(template)

    content = PadContent.model_validate_json(fake_llm())  # 형식이 틀리면 여기서 거부
    render(template, content, out)

    doc = Document(out)
    print("=== 렌더링 결과 ===")
    print("머리글:", doc.sections[0].header.paragraphs[0].text)
    for p in doc.paragraphs:
        print(f"[{p.style.name}] {p.text}")
    for row in doc.tables[0].rows:
        print(" | ".join(c.text for c in row.cells))

    before, after = set(formats(template)), set(formats(out))
    print(f"\n형식 속성(스타일·폰트·크기·굵기) 변화: {'없음 ✅' if after <= before else after - before}")
    leftover = [p.text for p in iter_paragraphs(doc) if "{{" in p.text]
    print(f"남은 placeholder: {leftover or '없음 ✅'}")

    print("\n=== 스키마가 막아주는 잘못된 LLM 출력 ===")
    try:
        PadContent.model_validate_json('{"part_number": "0412", "manufacturer": "ACME"}')
    except Exception as e:
        print(f"거부: {e.error_count()}개 오류 (부품번호 형식, 누락 필드)")


if __name__ == "__main__":
    main()
